"""Ground truth: manual awn lengths, cell-level presence, and instance polygons.

Three independent truths are used, all read-only from the research project:

* lengths   - ImageJ arc lengths per (line, strain, repeat, role) in PhenotypingResults.xlsx;
* cells     - human-confirmed presence of bottom/middle/top per page cell (validation);
* instances - reviewed awn / spikelet polygons. They are annotated on five overlapping
              column crops per page; here they are mapped back to page coordinates and
              de-duplicated, giving whole-awn instances usable as reconstruction truth.
"""
from __future__ import annotations

import json
import re
from dataclasses import dataclass
from functools import lru_cache

from openpyxl import load_workbook
from shapely import affinity
from shapely.geometry import Polygon
from shapely.geometry.base import BaseGeometry
from shapely.validation import make_valid

from .paths import paths

ROLES = ("bottom", "middle", "top")
POSITIVE_MM = 5.0  # phenotypic awn presence threshold defined by the project


def length_key(line: int, strain: int, repeat: int, role: str) -> str:
    return f"{int(line)}|{int(strain)}|{int(repeat)}|{role}"


@lru_cache(maxsize=1)
def manual_lengths() -> dict[str, float | None]:
    """workbook key -> raw manual length in mm (None for NA)."""
    book = load_workbook(paths().length_truth, read_only=True, data_only=True)
    rows = book["F3 Awn length long table"].iter_rows(values_only=True)
    header = tuple(next(rows)[:5])
    assert header == ("Line No.", "Strain No.", "Duplication", "Awn Type", "Awn Length (mm)"), header
    out: dict[str, float | None] = {}
    for values in rows:
        if not any(v is not None for v in values[:5]):
            continue
        line, strain, dup, role, length = values[:5]
        match = re.match(r"^dup\s*([123])$", str(dup).strip(), re.IGNORECASE)
        repeat = int(match.group(1)) if match else int(dup)
        if isinstance(length, str) and length.strip().casefold() in {"", "na", "n/a", "nan"}:
            length = None
        out[length_key(int(line), int(strain), repeat, str(role).strip().lower())] = (
            None if length is None else float(length)
        )
    book.close()
    return out


@lru_cache(maxsize=1)
def accession_mapping() -> dict[str, dict[int, tuple[int, int]]]:
    """page -> column -> (line, strain)."""
    raw = json.loads(paths().accession_mapping.read_text(encoding="utf-8"))["pages"]
    return {
        page: {int(c): (int(v[0]), int(v[1])) for c, v in cols.items() if v is not None}
        for page, cols in raw.items()
    }


@lru_cache(maxsize=1)
def cell_truth() -> dict[str, dict]:
    """"PAGE:C{c}R{r}" -> {present_roles, benchmark_include, ...} (validation pages)."""
    return json.loads(paths().validation_cell_truth.read_text(encoding="utf-8"))["cells"]


@dataclass(frozen=True)
class LengthTarget:
    page_id: str
    column: int
    repeat: int
    role: str
    key: str
    length_mm: float

    @property
    def positive(self) -> bool:
        return self.length_mm >= POSITIVE_MM


def length_targets(page_id: str) -> list[LengthTarget]:
    """Scored (cell, role) targets of one validation/test page.

    Only human-confirmed present roles of included cells with a numeric manual length
    enter the denominator, exactly as in the research project's evaluator.
    """
    cells, mapping, lengths = cell_truth(), accession_mapping(), manual_lengths()
    out = []
    for column in range(1, 6):
        for repeat in range(1, 4):
            truth = cells.get(f"{page_id}:C{column}R{repeat}")
            pair = mapping.get(page_id, {}).get(column)
            if truth is None or not truth["benchmark_include"] or pair is None:
                continue
            for role in truth["present_roles"]:
                key = length_key(*pair, repeat, role)
                value = lengths.get(key)
                if value is not None:
                    out.append(LengthTarget(page_id, column, repeat, role, key, value))
    return out


# --------------------------------------------------------------------------- instances


@dataclass(frozen=True)
class TruthInstance:
    id: str
    kind: str  # "awn" | "spikelet"
    geometry: BaseGeometry  # page pixels


@lru_cache(maxsize=1)
def _crop_manifest() -> dict[str, dict]:
    rows = [json.loads(line) for line in paths().crop_manifest.read_text(encoding="utf-8").splitlines() if line.strip()]
    return {row["crop"]: row for row in rows}


@lru_cache(maxsize=3)
def _coco(split: str) -> dict:
    return json.loads((paths().annotations / split / "_annotations.coco.json").read_text(encoding="utf-8"))


def column_ranges(page_id: str) -> dict[int, tuple[float, float]]:
    """Nominal x-range of the five accession columns on a page."""
    return {
        int(row["column_index"]): tuple(map(float, row["nominal_column_x_source"]))
        for row in _crop_manifest().values()
        if row["source"] == f"{page_id}.png"
    }


SAME_OBJECT_OVERLAP = 0.6  # intersection / smaller area for copies in different crops


def _polygonal(geometry):
    """Keep only the area part of a repaired polygon (drop stray lines / points)."""
    from shapely.ops import unary_union

    parts = [g for g in getattr(geometry, "geoms", [geometry]) if g.geom_type in ("Polygon", "MultiPolygon")]
    return unary_union(parts) if parts else geometry


@lru_cache(maxsize=256)
def _truth_instances(page_id: str, split: str) -> tuple[TruthInstance, ...]:
    coco = _coco(split)
    kinds = {c["id"]: c["name"] for c in coco["categories"]}
    images = {img["id"]: img for img in coco["images"] if img["extra"]["source_page"] == page_id}
    candidates = []  # (crop index, instance)
    for ann in coco["annotations"]:
        image = images.get(ann["image_id"])
        if image is None:
            continue
        crop = _crop_manifest()[image["file_name"]]
        coords = ann["segmentation"]
        if isinstance(coords, str):
            coords = json.loads(coords)
        ring = coords[0]
        polygon = _polygonal(make_valid(Polygon(list(zip(ring[0::2], ring[1::2])))))
        polygon = affinity.translate(polygon, xoff=float(crop["crop_xyxy"][0]))
        if not polygon.is_empty and polygon.area > 0:
            inst = TruthInstance(f"{page_id}:{ann['id']}", kinds[ann["category_id"]], polygon)
            candidates.append((int(crop["column_index"]), inst))

    # Copies of one object in neighbouring crops overlap strongly; keep the largest
    # (a copy cut by a crop edge is smaller). Objects within one crop are never merged.
    candidates.sort(key=lambda row: -row[1].geometry.area)
    kept: list[tuple[int, TruthInstance]] = []
    for crop, inst in candidates:
        duplicate = any(
            other_crop != crop
            and other.kind == inst.kind
            and other.geometry.intersects(inst.geometry)
            and other.geometry.intersection(inst.geometry).area / inst.geometry.area > SAME_OBJECT_OVERLAP
            for other_crop, other in kept
        )
        if not duplicate:
            kept.append((crop, inst))
    return tuple(inst for _, inst in kept)


def truth_instances(page_id: str, split: str) -> list[TruthInstance]:
    """Page-level reviewed instances, one copy per physical object.

    Each column crop overlaps its neighbours, so objects near column boundaries are
    annotated in two crops; those copies are merged here.
    """
    return list(_truth_instances(page_id, split))
